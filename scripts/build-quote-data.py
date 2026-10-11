"""Build public/quote-data.js (the instant-quote price table) from the Onyx_Distance_Price_Tables workbook.

Usage:  python3 scripts/build-quote-data.py path/to/Onyx_Distance_Price_Tables.xlsx [--endpoint https://script.google.com/.../exec]
Needs:  pip install openpyxl

The workbook lists 105 municipalities per airport/port. The website shows ~30 key zones instead:
each zone groups neighbouring municipalities and takes the HIGHEST price (and its distance/time)
among its members for every airport/port, so grouping never under-prices a trip.
Edit ZONES below to add, remove or regroup destinations.
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
HOURLY = {"rate": 100, "min": 3, "max": 12, "discFrom": 6, "discPct": 20}  # $/hour, hours, 20% off from 6 hours
RETURN_DISCOUNT_PCT = 20  # discount on the return ride of a round trip
ENDPOINT = ""  # Apps Script web app URL (docs/lead-capture/README.md); empty = requests open in WhatsApp

HUBS = [  # (sheet, id, label, kind)
    ("MIA", "MIA", "Miami International Airport (MIA)", "airport"),
    ("FLL", "FLL", "Fort Lauderdale Airport (FLL)", "airport"),
    ("PBI", "PBI", "Palm Beach Airport (PBI)", "airport"),
    ("PortMiami", "PORTMIA", "PortMiami Cruise Terminals", "port"),
    ("Port Everglades", "PORTEVER", "Port Everglades Cruise Terminals", "port"),
]

# (label shown to the customer, county, [municipalities from the workbook])
ZONES = [
    # Miami-Dade
    ("Downtown Miami / Brickell", "Miami-Dade", ["Miami (Downtown/Brickell)"]),
    ("Miami Beach / South Beach", "Miami-Dade", ["Miami Beach"]),
    ("Coral Gables", "Miami-Dade", ["Coral Gables"]),
    ("Key Biscayne", "Miami-Dade", ["Key Biscayne"]),
    ("Doral", "Miami-Dade", ["Doral"]),
    ("Miami Springs / Medley / Sweetwater", "Miami-Dade", ["Miami Springs", "Virginia Gardens", "West Miami", "Medley", "Sweetwater"]),
    ("Hialeah / Miami Lakes", "Miami-Dade", ["Hialeah", "Hialeah Gardens", "Miami Lakes"]),
    ("Kendall / Pinecrest / South Miami", "Miami-Dade", ["Kendall", "Pinecrest", "South Miami"]),
    ("Cutler Bay / Palmetto Bay", "Miami-Dade", ["Cutler Bay", "Palmetto Bay"]),
    ("Homestead / Florida City", "Miami-Dade", ["Homestead", "Florida City"]),
    ("North Miami / Miami Shores", "Miami-Dade", ["North Miami", "Miami Shores", "El Portal", "Biscayne Park"]),
    ("North Miami Beach / Miami Gardens", "Miami-Dade", ["North Miami Beach", "Miami Gardens", "Opa-locka"]),
    ("Bal Harbour / Surfside / Bay Harbor", "Miami-Dade", ["Bal Harbour", "Surfside", "Bay Harbor Islands", "Indian Creek"]),
    ("Sunny Isles Beach", "Miami-Dade", ["Sunny Isles Beach"]),
    ("Aventura / Golden Beach", "Miami-Dade", ["Aventura", "Golden Beach"]),
    # Broward
    ("Hollywood / Hallandale Beach / Dania Beach", "Broward", ["Hollywood", "Hallandale Beach", "Dania Beach", "Pembroke Park", "West Park"]),
    ("Miramar / Pembroke Pines", "Broward", ["Miramar", "Pembroke Pines"]),
    ("Davie / Weston / Cooper City", "Broward", ["Davie", "Weston", "Cooper City", "Southwest Ranches"]),
    ("Fort Lauderdale", "Broward", ["Fort Lauderdale"]),
    ("Plantation / Sunrise / Wilton Manors", "Broward", ["Plantation", "Sunrise", "Lauderhill", "Lauderdale Lakes", "Tamarac", "Oakland Park", "Wilton Manors"]),
    ("Pompano Beach / Lighthouse Point", "Broward", ["Pompano Beach", "Lauderdale-by-the-Sea", "Lighthouse Point"]),
    ("Coral Springs / Parkland", "Broward", ["Coral Springs", "Parkland", "Coconut Creek", "Margate"]),
    ("Deerfield Beach", "Broward", ["Deerfield Beach", "Hillsboro Beach"]),
    # Palm Beach
    ("Boca Raton", "Palm Beach", ["Boca Raton"]),
    ("Delray Beach / Highland Beach", "Palm Beach", ["Delray Beach", "Highland Beach"]),
    ("Boynton Beach / Ocean Ridge", "Palm Beach", ["Boynton Beach", "Ocean Ridge", "Gulf Stream", "Briny Breezes", "Golf"]),
    ("Lake Worth Beach / Lantana", "Palm Beach", ["Lake Worth Beach", "Lantana", "Hypoluxo", "Greenacres", "Palm Springs", "Atlantis", "Manalapan", "South Palm Beach"]),
    ("Wellington / Royal Palm Beach", "Palm Beach", ["Wellington", "Royal Palm Beach", "Loxahatchee Groves"]),
    ("West Palm Beach", "Palm Beach", ["West Palm Beach", "Glen Ridge", "Cloud Lake", "Haverhill", "Lake Clarke Shores"]),
    ("Palm Beach (island)", "Palm Beach", ["Palm Beach"]),
    ("Palm Beach Gardens / North Palm Beach", "Palm Beach", ["Palm Beach Gardens", "North Palm Beach", "Riviera Beach", "Lake Park", "Palm Beach Shores", "Mangonia Park", "Juno Beach"]),
    ("Jupiter / Tequesta", "Palm Beach", ["Jupiter", "Tequesta", "Jupiter Inlet Colony"]),
]

# Routes listed as static tables on rates.html: (title, hub id, [zone labels])
FEATURED = [
    ("From Miami International (MIA)", "MIA", ["Downtown Miami / Brickell", "Miami Beach / South Beach", "Coral Gables", "Key Biscayne", "Aventura / Golden Beach", "Sunny Isles Beach", "Bal Harbour / Surfside / Bay Harbor", "Fort Lauderdale", "Boca Raton", "Palm Beach (island)"]),
    ("From Fort Lauderdale (FLL)", "FLL", ["Fort Lauderdale", "Hollywood / Hallandale Beach / Dania Beach", "Aventura / Golden Beach", "Miami Beach / South Beach", "Downtown Miami / Brickell", "Boca Raton"]),
    ("PortMiami cruise terminals", "PORTMIA", ["Downtown Miami / Brickell", "Miami Beach / South Beach", "Aventura / Golden Beach", "Coral Gables"]),
    ("Port Everglades cruise terminals", "PORTEVER", ["Fort Lauderdale", "Hollywood / Hallandale Beach / Dania Beach", "Aventura / Golden Beach", "Boca Raton"]),
    ("From Palm Beach (PBI)", "PBI", ["Palm Beach (island)", "West Palm Beach", "Boca Raton", "Jupiter / Tequesta"]),
]


def main(xlsx):
    global ENDPOINT
    if "--endpoint" in sys.argv:
        ENDPOINT = sys.argv[sys.argv.index("--endpoint") + 1]
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    st = wb["Settings"]
    pricing = {"base": st["B5"].value, "perMile": st["B6"].value, "round": st["B7"].value}

    # raw[hub][municipality] = [miles, minutes, minutes_rush, price]
    raw = {}
    for sheet, hub_id, _, _ in HUBS:
        ws, t = wb[sheet], {}
        for r in range(5, ws.max_row + 1):
            name, price = ws.cell(r, 2).value, ws.cell(r, 7).value
            if not name or price is None:
                continue
            t[name] = [round(float(ws.cell(r, 4).value), 1), int(ws.cell(r, 5).value), int(ws.cell(r, 6).value), int(price)]
        raw[hub_id] = t

    places, tables = [], {h[1]: {} for h in HUBS}
    for i, (label, county, members) in enumerate(ZONES):
        places.append([label, county])
        for hub_id, t in raw.items():
            missing = [m for m in members if m not in t]
            if missing:
                raise SystemExit("Unknown municipality %s in zone %r (hub %s)" % (missing, label, hub_id))
            best = max((t[m] for m in members), key=lambda row: (row[3], row[0]))
            tables[hub_id][i] = best

    data = {
        "pricing": pricing,
        "hubs": [[h, label, kind] for _, h, label, kind in HUBS],
        "places": places,
        "table": tables,
        "lateNightPct": 30,
        "hourly": HOURLY,
        "returnDiscountPct": RETURN_DISCOUNT_PCT,
        "endpoint": ENDPOINT,
    }
    out = os.path.join(ROOT, "public", "quote-data.js")
    with open(out, "w", encoding="utf-8") as f:
        f.write("/* Generated by scripts/build-quote-data.py from the Onyx price workbook. Do not edit by hand. */\n")
        f.write("window.ONYX_QUOTE=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")
    print("zones:", len(places), "hubs:", len(tables), "->", os.path.normpath(out))


if __name__ == "__main__":
    main(sys.argv[1])
